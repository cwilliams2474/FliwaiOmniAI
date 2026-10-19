"""Document text extraction, chunking and retrieval (PostgreSQL full-text search)."""

import io
import logging
import re
import uuid
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from sqlalchemy import func, literal_column, select
from sqlalchemy.orm import Session

from app.models import DocStatus, Document, DocumentChunk, MemoryEntry

log = logging.getLogger("omniai.rag")

TEXT_TYPES = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".yaml", ".yml", ".ini", ".conf"}
SUPPORTED = TEXT_TYPES | {".pdf", ".docx", ".html", ".htm"}
CHUNK_CHARS = 1200
OVERLAP_CHARS = 150


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        if tag in ("p", "div", "br", "li", "h1", "h2", "h3", "h4", "tr"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def extract_pages(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    """Returns [(page_number_or_None, text)]. Raises ValueError for unsupported or unreadable files."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("This PDF is password protected.")
        pages = [(i + 1, (p.extract_text() or "")) for i, p in enumerate(reader.pages)]
        if not any(t.strip() for _, t in pages):
            raise ValueError("No text found in this PDF. It may be a scan; run it through OCR first.")
        return pages
    if ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return [(None, "\n\n".join(parts))]
    if ext in (".html", ".htm"):
        p = _HTMLText()
        p.feed(data.decode("utf-8", errors="replace"))
        return [(None, "".join(p.parts))]
    if ext in TEXT_TYPES:
        return [(None, data.decode("utf-8", errors="replace"))]
    raise ValueError(f"{ext or 'This file type'} isn't supported. Upload PDF, Word (.docx), HTML, Markdown or text files.")


def chunk_pages(pages: list[tuple[int | None, str]]) -> list[tuple[int | None, str]]:
    """Split into ~CHUNK_CHARS pieces on paragraph boundaries, with a little overlap for context."""
    chunks: list[tuple[int | None, str]] = []
    for page, raw in pages:
        txt = re.sub(r"[ \t\r\f\v]+", " ", raw)
        txt = re.sub(r"\n\s*\n+", "\n\n", txt).strip()
        if not txt:
            continue
        paras = [p.strip() for p in txt.split("\n\n") if p.strip()]
        buf = ""
        for para in paras:
            while len(para) > CHUNK_CHARS:  # very long paragraph: hard split on sentence or space
                cut = max(para.rfind(". ", 0, CHUNK_CHARS), para.rfind(" ", 0, CHUNK_CHARS))
                cut = cut + 1 if cut > CHUNK_CHARS // 2 else CHUNK_CHARS
                if buf:
                    chunks.append((page, buf))
                    buf = ""
                chunks.append((page, para[:cut].strip()))
                para = para[max(0, cut - OVERLAP_CHARS):].strip()
            if len(buf) + len(para) + 2 > CHUNK_CHARS and buf:
                chunks.append((page, buf))
                buf = buf[-OVERLAP_CHARS:].split(" ", 1)[-1] + "\n\n" + para
            else:
                buf = f"{buf}\n\n{para}" if buf else para
        if buf:
            chunks.append((page, buf))
    return chunks


def index_document(db: Session, doc: Document, data: bytes) -> None:
    try:
        pages = extract_pages(doc.filename, data)
        chunks = chunk_pages(pages)
        if not chunks:
            raise ValueError("The file has no readable text.")
        db.query(DocumentChunk).filter(DocumentChunk.document_id == doc.id).delete()
        for i, (page, content) in enumerate(chunks):
            db.add(DocumentChunk(document_id=doc.id, org_id=doc.org_id, ordinal=i, page=page, content=content))
        doc.chunk_count = len(chunks)
        doc.status = DocStatus.indexed
        doc.error = ""
    except Exception as e:  # noqa: BLE001 - any parser failure becomes a visible document error
        log.warning("Indexing %s failed: %s", doc.filename, e)
        doc.status = DocStatus.failed
        doc.error = str(e)[:1000] if isinstance(e, ValueError) else "The file couldn't be read. It may be damaged."
        doc.chunk_count = 0
    db.commit()


_WORD = re.compile(r"[A-Za-z0-9][A-Za-z0-9._-]*")


def to_or_query(q: str) -> str | None:
    """Turn free text into an OR'ed tsquery string of safe terms, e.g. 'vpn | dns | restart'."""
    words = [w.strip("._-").lower() for w in _WORD.findall(q)]
    words = [w for w in words if len(w) > 1][:24]
    if not words:
        return None
    return " | ".join(dict.fromkeys(words))


@dataclass
class Hit:
    chunk_id: int
    document_id: uuid.UUID
    filename: str
    page: int | None
    content: str
    score: float


def search(db: Session, org_id: uuid.UUID, q: str, limit: int = 6) -> list[Hit]:
    tq = to_or_query(q)
    if not tq:
        return []
    query = func.to_tsquery(literal_column("'english'::regconfig"), tq)
    rank = func.ts_rank_cd(DocumentChunk.tsv, query, 32).label("score")
    rows = db.execute(
        select(DocumentChunk.id, DocumentChunk.document_id, Document.filename, DocumentChunk.page, DocumentChunk.content, rank)
        .join(Document, Document.id == DocumentChunk.document_id)
        .where(DocumentChunk.org_id == org_id, Document.status == DocStatus.indexed, DocumentChunk.tsv.op("@@")(query))
        .order_by(rank.desc())
        .limit(limit)
    ).all()
    return [Hit(chunk_id=r[0], document_id=r[1], filename=r[2], page=r[3], content=r[4], score=float(r[5])) for r in rows]


def org_memory(db: Session, org_id: uuid.UUID, limit: int = 40) -> list[MemoryEntry]:
    return list(db.scalars(select(MemoryEntry).where(MemoryEntry.org_id == org_id).order_by(MemoryEntry.updated_at.desc()).limit(limit)))
